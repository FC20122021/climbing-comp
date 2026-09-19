# -*- coding: utf-8 -*-
"""
攀岩比赛系统 - 报名 + 打分 + 权限区分
数据持久化：优先使用 Supabase PostgreSQL（通过 DATABASE_URL 环境变量）
本地开发：自动回退到 SQLite
支持中英双语（首页、报名页、登录页、裁判打分页）
"""

import os
from datetime import datetime, timedelta
from io import BytesIO
from functools import wraps

from flask import Flask, render_template, request, redirect, url_for, session, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from openpyxl import Workbook

# -------------------- 应用初始化 --------------------
app = Flask(__name__)

# 数据库配置：优先使用环境变量中的 DATABASE_URL（Render + Supabase）
database_url = os.environ.get('DATABASE_URL')
if database_url:
    # 兼容 postgres:// 与 postgresql://
    if database_url.startswith('postgres://'):
        database_url = database_url.replace('postgres://', 'postgresql://', 1)
    # Supabase 要求 SSL 连接
    if 'sslmode' not in database_url:
        separator = '&' if '?' in database_url else '?'
        database_url = database_url + separator + 'sslmode=require'
    app.config['SQLALCHEMY_DATABASE_URI'] = database_url
else:
    # 本地开发使用 SQLite
    db_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'instance', 'climbing.db')
    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    app.config['SQLALCHEMY_DATABASE_URI'] = f'sqlite:///{db_path}'

app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SECRET_KEY'] = 'your-secret-key-here'

db = SQLAlchemy(app)


# -------------------- 数据模型 --------------------
class Athlete(db.Model):
    __tablename__ = 'athlete'
    id = db.Column(db.Integer, primary_key=True)
    number = db.Column(db.Integer, unique=True, nullable=True)
    name = db.Column(db.String(50), nullable=False, unique=True)
    phone = db.Column(db.String(20), nullable=False, unique=True)
    level = db.Column(db.String(10))
    card_type = db.Column(db.String(20))
    climbing_years = db.Column(db.String(20))
    climbing_frequency = db.Column(db.String(20))
    climbing_days = db.Column(db.String(50))
    discovery_channel = db.Column(db.String(50))
    competition_channel = db.Column(db.String(50))
    created_at = db.Column(db.DateTime, default=datetime.now)


class ScoreLog(db.Model):
    __tablename__ = 'score_log'
    id = db.Column(db.Integer, primary_key=True)
    athlete_id = db.Column(db.Integer, db.ForeignKey('athlete.id'), nullable=False)
    route_id = db.Column(db.Integer, nullable=False)
    result = db.Column(db.String(10), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.now)
    status = db.Column(db.String(10), default='valid')
    judge_id = db.Column(db.Integer, db.ForeignKey('judge.id'), nullable=True)

    athlete = db.relationship('Athlete', backref=db.backref('scores', lazy=True))


class Route(db.Model):
    __tablename__ = 'route'
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(20), nullable=False)


class Judge(db.Model):
    __tablename__ = 'judge'
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(50), unique=True, nullable=False)
    password = db.Column(db.String(100), nullable=False)
    role = db.Column(db.String(20), default='judge')


class CompetitionState(db.Model):
    __tablename__ = 'competition_state'
    id = db.Column(db.Integer, primary_key=True)
    status = db.Column(db.String(20), default='not_started')
    start_time = db.Column(db.DateTime, nullable=True)
    end_time = db.Column(db.DateTime, nullable=True)
    buffer_end_time = db.Column(db.DateTime, nullable=True)


# -------------------- 启动时初始化数据库 --------------------
with app.app_context():
    try:
        db.create_all()

        if Judge.query.count() == 0:
            db.session.add(Judge(username='admin', password='admin123', role='admin'))
            db.session.commit()
        else:
            admin = Judge.query.filter_by(username='admin').first()
            if admin and admin.role != 'admin':
                admin.role = 'admin'
                db.session.commit()

        if Route.query.count() == 0:
            db.session.add_all([Route(name=f'线路{i}') for i in range(1, 9)])
            db.session.commit()

        if CompetitionState.query.count() == 0:
            db.session.add(CompetitionState(status='not_started'))
            db.session.commit()
    except Exception as e:
        print(f'[init_db] {e}')
        db.session.rollback()


def init_db():
    with app.app_context():
        db.create_all()


# -------------------- 辅助函数 --------------------
def get_competition_state():
    return CompetitionState.query.first()


def update_competition_status_if_needed():
    state = get_competition_state()
    if not state:
        return state

    now = datetime.now()
    if state.status == 'running' and state.end_time and now >= state.end_time:
        state.status = 'buffer'
        state.buffer_end_time = now + timedelta(minutes=1)
        db.session.commit()
    elif state.status == 'buffer' and state.buffer_end_time and now >= state.buffer_end_time:
        state.status = 'ended'
        db.session.commit()
    return state


def get_athlete_final_scores(athlete_id):
    logs = ScoreLog.query.filter_by(athlete_id=athlete_id, status='valid').order_by(ScoreLog.created_at.asc()).all()

    route_logs = {}
    for log in logs:
        if log.route_id not in route_logs:
            route_logs[log.route_id] = []
        route_logs[log.route_id].append(log)

    total_score = 0.0
    top_count = 0
    zone_count = 0
    total_fail_attempts = 0
    per_route = {}

    for route_id, log_list in route_logs.items():
        has_top = any(log.result == 'top' for log in log_list)
        has_zone = any(log.result == 'zone' for log in log_list)
        final_result = 'fail'
        if has_top:
            final_result = 'top'
        elif has_zone:
            final_result = 'zone'

        penalty_attempts = 0
        if final_result == 'top':
            first_top_index = next(i for i, log in enumerate(log_list) if log.result == 'top')
            penalty_attempts = first_top_index
        elif final_result == 'zone':
            first_zone_index = next(i for i, log in enumerate(log_list) if log.result == 'zone')
            penalty_attempts = first_zone_index

        base_score = 0
        if final_result == 'top':
            base_score = 25
            top_count += 1
        elif final_result == 'zone':
            base_score = 10
            zone_count += 1

        deduction = penalty_attempts * 0.1
        route_score = base_score - deduction

        total_score += route_score

        fail_count = sum(1 for log in log_list if log.result == 'fail')
        total_fail_attempts += fail_count

        per_route[route_id] = {
            'final_result': final_result,
            'score': route_score,
            'penalty_attempts': penalty_attempts,
            'attempts': len(log_list)
        }

    return {
        'total_score': round(total_score, 1),
        'top_count': top_count,
        'zone_count': zone_count,
        'total_fail_attempts': total_fail_attempts,
        'per_route': per_route
    }


def get_leaderboard_data():
    athletes = Athlete.query.all()
    leaderboard = []
    for athlete in athletes:
        has_scores = ScoreLog.query.filter_by(athlete_id=athlete.id, status='valid').first()
        if not has_scores:
            continue
        stats = get_athlete_final_scores(athlete.id)
        leaderboard.append({
            'athlete_id': athlete.id,
            'number': athlete.number,
            'name': athlete.name,
            'total_score': stats['total_score'],
            'top_count': stats['top_count'],
            'zone_count': stats['zone_count'],
            'total_fail_attempts': stats['total_fail_attempts']
        })

    leaderboard.sort(key=lambda x: (-x['total_score'], -x['top_count'], -x['zone_count'], x['total_fail_attempts']))
    for idx, item in enumerate(leaderboard):
        item['rank'] = idx + 1
    return leaderboard


# -------------------- 装饰器 --------------------
def judge_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'judge_id' not in session:
            return redirect(url_for('judge_login'))
        return f(*args, **kwargs)
    return decorated_function


def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'judge_id' not in session:
            return redirect(url_for('judge_login'))
        role = session.get('judge_role')
        if role != 'admin':
            return redirect(url_for('judge_dashboard'))
        return f(*args, **kwargs)
    return decorated_function


# -------------------- 语言切换 --------------------
@app.route('/set_language/<lang>')
def set_language(lang):
    if lang in ['zh', 'en']:
        session['lang'] = lang
    return redirect(request.referrer or url_for('index'))


# -------------------- 首页入口 --------------------
@app.route('/')
def index():
    if session.get('lang') == 'en':
        return render_template('public/index_en.html')
    return render_template('public/index.html')


# -------------------- 比赛名单（公开） --------------------
@app.route('/athletes')
def athletes_list():
    all_athletes = Athlete.query.all()

    with_number = [a for a in all_athletes if a.number is not None]
    with_number.sort(key=lambda x: x.number)

    without_number = [a for a in all_athletes if a.number is None]
    without_number.sort(key=lambda x: x.name.lower())

    sorted_athletes = with_number + without_number

    if session.get('lang') == 'en':
        return render_template('public/athletes_en.html', athletes=sorted_athletes)
    return render_template('public/athletes.html', athletes=sorted_athletes)


# -------------------- 历史比赛（公开） --------------------
@app.route('/competitions')
def competitions():
    return render_template('public/competitions.html')


# -------------------- 报名（公开） --------------------
@app.route('/signup', methods=['GET'])
def signup():
    message = request.args.get('message', '')
    msg_type = request.args.get('type', 'success')
    if session.get('lang') == 'en':
        return render_template('public/signup_en.html', message=message, msg_type=msg_type)
    return render_template('public/signup.html', message=message, msg_type=msg_type)


@app.route('/signup', methods=['POST'])
def signup_post():
    name = request.form.get('name', '').strip()
    country_code = request.form.get('country_code', '+86').strip()
    phone = request.form.get('phone', '').strip()
    level = request.form.get('level', '').strip()
    card_type = request.form.get('card_type', '').strip()
    climbing_years = request.form.get('climbing_years', '').strip()
    climbing_frequency = request.form.get('climbing_frequency', '').strip()
    climbing_days_list = request.form.getlist('climbing_days')
    climbing_days = ','.join(climbing_days_list)
    discovery_channel = request.form.get('discovery_channel', '').strip()
    competition_channel = request.form.get('competition_channel', '').strip()
    number_str = request.form.get('number', '').strip()
    agree = request.form.get('agree')

    is_en = session.get('lang') == 'en'

    if not name:
        return redirect(url_for('signup', message='Name is required' if is_en else '姓名不能为空', type='error'))

    if not phone:
        return redirect(url_for('signup', message='Phone is required' if is_en else '手机号不能为空', type='error'))

    if country_code == '+86':
        if len(phone) != 11 or not phone.isdigit():
            return redirect(url_for('signup', message='Mainland China phone must be 11 digits' if is_en else '中国大陆手机号应为11位数字', type='error'))
    elif country_code == '+852':
        if len(phone) != 8 or not phone.isdigit():
            return redirect(url_for('signup', message='Hong Kong phone must be 8 digits' if is_en else '香港手机号应为8位数字', type='error'))
    else:
        return redirect(url_for('signup', message='Unsupported country code' if is_en else '暂不支持该国家代码', type='error'))

    if not agree:
        return redirect(url_for('signup', message='Please agree to the rules and pledge' if is_en else '请阅读并同意比赛守则和承诺书后再报名', type='error'))

    if not card_type:
        return redirect(url_for('signup', message='Please select a membership type' if is_en else '请选择会员卡类型，本次比赛仅限会员参加', type='error'))

    number = None
    if number_str:
        try:
            number = int(number_str)
        except ValueError:
            return redirect(url_for('signup', message='Number must be a digit' if is_en else '比赛编号必须是数字', type='error'))
        if number < 0 or number > 999:
            return redirect(url_for('signup', message='Number must be between 0 and 999' if is_en else '比赛编号必须在 0 到 999 之间', type='error'))
        existing_number = Athlete.query.filter_by(number=number).first()
        if existing_number:
            return redirect(url_for('signup', message=f'Number {number} is taken' if is_en else f'编号 {number} 已被占用，请选择其他编号', type='error'))

    full_phone = f'{country_code}{phone}'

    existing_name = Athlete.query.filter_by(name=name).first()
    if existing_name:
        return redirect(url_for('signup', message='This name is already used' if is_en else '该姓名/昵称已被使用', type='error'))

    existing_phone = Athlete.query.filter_by(phone=full_phone).first()
    if existing_phone:
        return redirect(url_for('signup', message='This phone is already registered' if is_en else '该手机号已报名', type='error'))

    new_athlete = Athlete(
        name=name,
        phone=full_phone,
        number=number,
        level=level if level else None,
        card_type=card_type if card_type else None,
        climbing_years=climbing_years if climbing_years else None,
        climbing_frequency=climbing_frequency if climbing_frequency else None,
        climbing_days=climbing_days if climbing_days else None,
        discovery_channel=discovery_channel if discovery_channel else None,
        competition_channel=competition_channel if competition_channel else None,
    )
    db.session.add(new_athlete)
    try:
        db.session.commit()
    except:
        db.session.rollback()
        return redirect(url_for('signup', message='Registration failed' if is_en else '报名失败，请检查信息是否重复', type='error'))

    return redirect(url_for('poster'))


# -------------------- API：查询编号是否被占用 --------------------
@app.route('/api/check_number/<int:number>')
def api_check_number(number):
    if number < 0 or number > 999:
        return jsonify({'taken': True, 'error': 'Number out of range'})
    athlete = Athlete.query.filter_by(number=number).first()
    return jsonify({'taken': athlete is not None})


# -------------------- 报名成功海报页 --------------------
@app.route('/poster')
def poster():
    return render_template('public/poster.html')


# -------------------- 裁判登录/登出 --------------------
@app.route('/judge/login', methods=['GET'])
def judge_login():
    if session.get('lang') == 'en':
        return render_template('auth/judge_login_en.html', error=None)
    return render_template('auth/judge_login.html', error=None)


@app.route('/judge/login', methods=['POST'])
def judge_login_post():
    username = request.form.get('username', '').strip()
    password = request.form.get('password', '').strip()
    judge = Judge.query.filter_by(username=username).first()
    if judge and judge.password == password:
        session['judge_id'] = judge.id
        session['judge_username'] = judge.username
        session['judge_role'] = judge.role
        if judge.role == 'admin':
            return redirect(url_for('admin_control'))
        else:
            return redirect(url_for('judge_dashboard'))
    if session.get('lang') == 'en':
        return render_template('auth/judge_login_en.html', error='Invalid username or password')
    return render_template('auth/judge_login.html', error='用户名或密码错误')


@app.route('/judge/logout')
def judge_logout():
    session.pop('judge_id', None)
    session.pop('judge_username', None)
    session.pop('judge_role', None)
    return redirect(url_for('index'))


# -------------------- 裁判打分页面 --------------------
@app.route('/judge')
@judge_required
def judge_dashboard():
    routes = Route.query.all()
    if session.get('lang') == 'en':
        return render_template('judge/judge_en.html', routes=routes, judge_username=session.get('judge_username'))
    return render_template('judge/judge.html', routes=routes, judge_username=session.get('judge_username'))


# -------------------- API：查询选手 --------------------
@app.route('/api/athlete/<int:number>')
@judge_required
def api_get_athlete(number):
    athlete = Athlete.query.filter_by(number=number).first()
    if athlete:
        return jsonify({'name': athlete.name, 'id': athlete.id})
    return jsonify({'error': '选手不存在'}), 404


# -------------------- API：提交成绩 --------------------
@app.route('/api/submit_score', methods=['POST'])
@judge_required
def api_submit_score():
    state = update_competition_status_if_needed()
    if not state or state.status not in ['running', 'buffer']:
        return jsonify({'error': '比赛尚未开始或已结束，无法提交成绩'}), 400

    data = request.get_json()
    route_id = data.get('route_id')
    result = data.get('result')

    judge_username = session.get('judge_username', '')
    if judge_username.startswith('judge') and judge_username[5:].isdigit():
        athlete_number = int(judge_username[5:])
    else:
        athlete_number = data.get('athlete_number')

    if not athlete_number or not route_id or result not in ['fail', 'zone', 'top']:
        return jsonify({'error': '参数错误'}), 400

    athlete = Athlete.query.filter_by(number=athlete_number).first()
    if not athlete:
        return jsonify({'error': '选手不存在'}), 404

    route = Route.query.get(route_id)
    if not route:
        return jsonify({'error': '线路不存在'}), 404

    existing_top = ScoreLog.query.filter_by(
        athlete_id=athlete.id,
        route_id=route_id,
        result='top',
        status='valid'
    ).first()
    if existing_top:
        return jsonify({'error': '该选手已在该线路完攀，不能重复提交'}), 400

    log = ScoreLog(
        athlete_id=athlete.id,
        route_id=route_id,
        result=result,
        judge_id=session['judge_id']
    )
    db.session.add(log)
    db.session.commit()

    return jsonify({'success': True, 'message': '成绩已记录'})


# -------------------- API：撤销上一步 --------------------
@app.route('/api/undo_last', methods=['POST'])
@judge_required
def api_undo_last():
    judge_id = session['judge_id']
    last_log = ScoreLog.query.filter_by(judge_id=judge_id, status='valid').order_by(ScoreLog.created_at.desc()).first()
    if not last_log:
        return jsonify({'error': '没有可撤销的记录'}), 404

    last_log.status = 'revoked'
    db.session.commit()
    return jsonify({'success': True, 'message': '已撤销'})


# -------------------- API：实时排名（管理员） --------------------
@app.route('/api/leaderboard')
@admin_required
def api_leaderboard():
    leaderboard = get_leaderboard_data()
    return jsonify(leaderboard)


# -------------------- API：比赛状态（管理员） --------------------
@app.route('/api/competition_status')
@admin_required
def api_competition_status():
    state = update_competition_status_if_needed()
    if not state:
        return jsonify({'status': 'not_started', 'remaining_seconds': 0})

    remaining = 0
    if state.status == 'running' and state.end_time:
        remaining = max(0, int((state.end_time - datetime.now()).total_seconds()))
    elif state.status == 'buffer' and state.buffer_end_time:
        remaining = max(0, int((state.buffer_end_time - datetime.now()).total_seconds()))

    return jsonify({
        'status': state.status,
        'remaining_seconds': remaining,
        'start_time': state.start_time.strftime('%Y-%m-%d %H:%M:%S') if state.start_time else None,
        'end_time': state.end_time.strftime('%Y-%m-%d %H:%M:%S') if state.end_time else None
    })


# -------------------- 大屏排名（管理员） --------------------
@app.route('/leaderboard')
@admin_required
def leaderboard_page():
    return render_template('leaderboard.html')


# -------------------- 抽奖转盘（管理员） --------------------
@app.route('/admin/lottery')
@admin_required
def admin_lottery():
    # 只取有编号的选手，按编号升序
    athletes = Athlete.query.filter(Athlete.number.isnot(None)).order_by(Athlete.number).all()
    return render_template('admin/admin_lottery.html', athletes=athletes)


# -------------------- 管理后台 --------------------
@app.route('/admin/routes', methods=['GET'])
@admin_required
def admin_routes():
    routes = Route.query.all()
    return render_template('admin/admin_routes.html', routes=routes)


@app.route('/admin/routes', methods=['POST'])
@admin_required
def admin_routes_post():
    name = request.form.get('name', '').strip()
    if name:
        route = Route(name=name)
        db.session.add(route)
        db.session.commit()
    return redirect(url_for('admin_routes'))


@app.route('/admin/signups')
@admin_required
def admin_signups():
    athletes = Athlete.query.order_by(Athlete.created_at.desc()).all()
    return render_template('admin/admin_signups.html', athletes=athletes)


@app.route('/admin/export_signups')
@admin_required
def export_signups():
    athletes = Athlete.query.order_by(Athlete.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "报名记录"

    headers = ['编号', '姓名', '手机号', '水平', '会员卡类型', '攀岩年限', '攀岩频率', '常来日', '了解渠道', '比赛渠道', '报名时间']
    ws.append(headers)

    for athlete in athletes:
        ws.append([
            athlete.number if athlete.number is not None else '',
            athlete.name,
            athlete.phone,
            athlete.level or '',
            athlete.card_type or '',
            athlete.climbing_years or '',
            athlete.climbing_frequency or '',
            athlete.climbing_days or '',
            athlete.discovery_channel or '',
            athlete.competition_channel or '',
            athlete.created_at.strftime('%Y-%m-%d %H:%M:%S') if athlete.created_at else ''
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='signups.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/admin/import', methods=['GET'])
@admin_required
def admin_import():
    return render_template('admin/admin_import.html')


@app.route('/admin/import', methods=['POST'])
@admin_required
def admin_import_post():
    file = request.files.get('file')
    if not file:
        return redirect(url_for('admin_import'))

    from openpyxl import load_workbook
    wb = load_workbook(file)
    ws = wb.active

    updated = 0
    skipped = 0
    errors = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or len(row) < 2:
            continue

        raw_number = row[0]
        raw_name = str(row[1]).strip() if row[1] is not None else ''

        if raw_number is None or raw_name == '':
            errors.append(f"第{row_idx}行：编号或姓名为空，已跳过")
            skipped += 1
            continue

        try:
            number = int(raw_number)
        except (ValueError, TypeError):
            errors.append(f"第{row_idx}行：编号 '{raw_number}' 不是有效数字，已跳过")
            skipped += 1
            continue

        existing_number = Athlete.query.filter(Athlete.number == number, Athlete.name != raw_name).first()
        if existing_number:
            errors.append(f"第{row_idx}行：编号 {number} 已被选手 {existing_number.name} 使用，已跳过")
            skipped += 1
            continue

        athlete = Athlete.query.filter_by(name=raw_name).first()
        if athlete:
            athlete.number = number
            updated += 1
        else:
            errors.append(f"第{row_idx}行：姓名 '{raw_name}' 未找到，已跳过")
            skipped += 1

    db.session.commit()
    message = f'导入完成：更新 {updated} 条，跳过 {skipped} 条'
    return render_template('admin/admin_import.html', message=message, errors=errors)


@app.route('/admin/control')
@admin_required
def admin_control():
    message = request.args.get('message', '')
    msg_type = request.args.get('type', 'success')
    return render_template('admin/admin_control.html', message=message, msg_type=msg_type)



@app.route('/admin/start', methods=['POST'])
@admin_required
def admin_start():
    state = get_competition_state()
    if not state:
        state = CompetitionState()
        db.session.add(state)

    if state.status in ['running', 'buffer']:
        return redirect(url_for('admin_control', message='比赛已经开始了，不能重复开始', type='error'))

    state.status = 'running'
    state.start_time = datetime.now()
    state.end_time = datetime.now() + timedelta(hours=2)
    state.buffer_end_time = None
    db.session.commit()
    return redirect(url_for('leaderboard_page'))


@app.route('/admin/create_judge_accounts', methods=['POST'])
@admin_required
def create_judge_accounts():
    password = 'ale26666'
    created = 0
    skipped = 0
    athletes = Athlete.query.filter(Athlete.number.isnot(None)).all()
    for athlete in athletes:
        username = f'judge{athlete.number}'
        if Judge.query.filter_by(username=username).first():
            skipped += 1
            continue
        judge = Judge(username=username, password=password, role='judge')
        db.session.add(judge)
        created += 1
    db.session.commit()
    msg = f'已创建 {created} 个账号，跳过 {skipped} 个已存在的'
    return redirect(url_for('admin_control', message=msg, type='success'))


@app.route('/admin/export_ranking')
@admin_required
def export_ranking():
    leaderboard = get_leaderboard_data()
    wb = Workbook()
    ws = wb.active
    ws.title = "总排名"
    headers = ['排名', '选手编号', '姓名', '总积分', 'Top数', 'Zone数']
    ws.append(headers)
    for item in leaderboard:
        ws.append([item['rank'], item['number'], item['name'], item['total_score'], item['top_count'], item['zone_count']])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='ranking.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/admin/export_route_details')
@admin_required
def export_route_details():
    athletes = Athlete.query.all()
    routes = Route.query.all()
    wb = Workbook()
    ws = wb.active
    ws.title = "线路明细"
    headers = ['选手编号', '姓名'] + [route.name for route in routes]
    ws.append(headers)

    for athlete in athletes:
        stats = get_athlete_final_scores(athlete.id)
        row = [athlete.number, athlete.name]
        for route in routes:
            route_info = stats['per_route'].get(route.id)
            if route_info:
                row.append(f"{route_info['final_result']} ({route_info['score']})")
            else:
                row.append('')
        ws.append(row)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='route_details.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/admin/export_logs')
@admin_required
def export_logs():
    logs = ScoreLog.query.order_by(ScoreLog.created_at.asc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "操作日志"
    headers = ['ID', '选手编号', '姓名', '线路', '结果', '裁判', '时间', '状态']
    ws.append(headers)
    for log in logs:
        athlete = Athlete.query.get(log.athlete_id)
        route = Route.query.get(log.route_id)
        judge = Judge.query.get(log.judge_id) if log.judge_id else None
        ws.append([
            log.id,
            athlete.number if athlete else '',
            athlete.name if athlete else '',
            route.name if route else '',
            log.result,
            judge.username if judge else '',
            log.created_at.strftime('%Y-%m-%d %H:%M:%S') if log.created_at else '',
            log.status
        ])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='logs.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# -------------------- 启动 --------------------
if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
